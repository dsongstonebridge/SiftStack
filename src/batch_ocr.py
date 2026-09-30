"""Batch OCR for scanned foreclosure petition PDFs.

Petitions arrive as scans with no text layer, so every page has to go through
Tesseract before anyone can read it. This OCRs every page of every PDF in a
folder in parallel (one process per core), caches the text per PDF, and prints
a case-number / file-date summary flagged against cases already processed.

    python src/batch_ocr.py                              # Desktop\\Foreclosure pdfs
    python src/batch_ocr.py --folder "<dir>" --workers 8
    python src/batch_ocr.py --force                      # re-OCR cached PDFs

Output (gitignored, under output/petition_ocr/):
    <stem>.txt       every page, "=== PAGE n/N ===" separated (body: 300dpi psm 4)
    <stem>.hdr.txt   page 1 re-read at 500dpi psm 6 + psm 11 for the stamped
                     case number and file date
    summary.csv      one row per PDF: case number, file date, already processed?

Free and local: nothing here touches DataSift or any billed service.
"""

import argparse
import csv
import glob
import os
import re
import sys
import time
from concurrent.futures import ProcessPoolExecutor, as_completed

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import pypdfium2 as pdfium  # noqa: E402

from image_utils import ocr_page  # noqa: E402

DEFAULT_FOLDER = os.path.join(os.path.expanduser("~"), "OneDrive", "Desktop", "Foreclosure pdfs")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DIR = os.path.join(ROOT, "output", "petition_ocr")

BODY_DPI = 300
HEADER_DPI = 500

CASE_RX = re.compile(r"\b[CcGg][JjI1l]\s*[-– ]\s*(20\d\d)\s*[-– ]\s*0*(\d{2,6})\b")
MONTHS = ("JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC")
DATE_RX = re.compile(
    r"\b(JAN|FEB|MAR|APR|MAY|JUN|JUL|AUG|SEPT?|OCT|NOV|DEC)[A-Z]*\.?\s+(\d{1,2}),?\s+(20\d\d)\b",
    re.IGNORECASE,
)


def norm_case(year: str, num: str) -> str:
    """CJ-2026-04287 and CJ-2026-4287 are the same case: drop leading zeros."""
    return f"CJ-{year}-{int(num)}"


def _ocr_task(pdf_path: str, index: int, header: bool) -> tuple:
    """Render one page and OCR it. Runs in a worker process."""
    pdf = pdfium.PdfDocument(pdf_path)
    try:
        page = pdf[index]
        if header:
            img = page.render(scale=HEADER_DPI / 72).to_pil()
            text = ("--- psm 6 ---\n" + ocr_page(img, psm=6)
                    + "\n--- psm 11 ---\n" + ocr_page(img, psm=11))
        else:
            img = page.render(scale=BODY_DPI / 72).to_pil()
            text = ocr_page(img, psm=4)
    finally:
        pdf.close()
    return pdf_path, index, header, text


def known_cases() -> set:
    """Case numbers already on a processed petition sheet in output/."""
    try:
        import openpyxl
    except ImportError:
        return set()
    seen = set()
    for path in glob.glob(os.path.join(ROOT, "output", "petition_batch*.xlsx")):
        try:
            ws = openpyxl.load_workbook(path, read_only=True).active
            rows = ws.iter_rows(values_only=True)
            header = next(rows)
            if "Case Number" not in header:
                continue
            col = header.index("Case Number")
            for row in rows:
                m = CASE_RX.search(str(row[col] or ""))
                if m:
                    seen.add(norm_case(*m.groups()))
        except Exception as exc:  # a locked or odd sheet should not stop the OCR
            print(f"WARN could not read {os.path.basename(path)}: {exc}")
    return seen


def find_case(text: str) -> str:
    m = CASE_RX.search(text)
    return norm_case(*m.groups()) if m else ""


def find_filed(text: str) -> str:
    """First plausible stamped date near a FILED marker, else the first date."""
    pool = text
    idx = text.upper().find("FILED")
    if idx >= 0:
        pool = text[idx: idx + 400] + "\n" + text
    for m in DATE_RX.finditer(pool):
        mon = MONTHS.index(m.group(1).upper()[:3]) + 1
        return f"{m.group(3)}-{mon:02d}-{int(m.group(2)):02d}"
    return ""


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--folder", default=DEFAULT_FOLDER)
    ap.add_argument("--workers", type=int, default=os.cpu_count() or 4)
    ap.add_argument("--force", action="store_true", help="re-OCR PDFs that already have cached text")
    args = ap.parse_args()

    os.makedirs(OUT_DIR, exist_ok=True)
    pdfs = sorted(glob.glob(os.path.join(args.folder, "*.pdf")))
    if not pdfs:
        print(f"No PDFs in {args.folder}")
        return 1

    def stem(p):
        return os.path.splitext(os.path.basename(p))[0]

    todo = [p for p in pdfs if args.force or not os.path.exists(os.path.join(OUT_DIR, stem(p) + ".txt"))]
    counts = {}
    for p in todo:
        doc = pdfium.PdfDocument(p)
        counts[p] = len(doc)
        doc.close()
    tasks = [(p, i, False) for p in todo for i in range(counts[p])] + [(p, 0, True) for p in todo]
    print(f"{len(pdfs)} PDFs, {len(pdfs) - len(todo)} cached, OCR'ing {len(todo)} "
          f"({sum(counts.values())} pages) on {args.workers} workers")

    pages, headers = {}, {}
    start, done = time.time(), 0
    if tasks:
        with ProcessPoolExecutor(max_workers=args.workers) as ex:
            futures = [ex.submit(_ocr_task, *t) for t in tasks]
            for fut in as_completed(futures):
                path, index, header, text = fut.result()
                if header:
                    headers[path] = text
                else:
                    pages.setdefault(path, {})[index] = text
                done += 1
                if done % 50 == 0 or done == len(tasks):
                    print(f"  {done}/{len(tasks)} pages  {time.time() - start:.0f}s", flush=True)

    for p in todo:
        n = counts[p]
        body = "".join(f"\n=== PAGE {i + 1}/{n} ===\n{pages[p][i]}" for i in range(n))
        with open(os.path.join(OUT_DIR, stem(p) + ".txt"), "w", encoding="utf-8") as f:
            f.write(body)
        with open(os.path.join(OUT_DIR, stem(p) + ".hdr.txt"), "w", encoding="utf-8") as f:
            f.write(headers[p])

    seen = known_cases()
    rows, by_case = [], {}
    for p in pdfs:
        with open(os.path.join(OUT_DIR, stem(p) + ".hdr.txt"), encoding="utf-8") as f:
            hdr = f.read()
        with open(os.path.join(OUT_DIR, stem(p) + ".txt"), encoding="utf-8") as f:
            body = f.read()
        first = body.split("=== PAGE 2/", 1)[0]
        case = find_case(hdr) or find_case(first)
        filed = find_filed(hdr) or find_filed(first)
        status = "ALREADY PROCESSED" if case and case in seen else ("NO CASE # FOUND" if not case else "new")
        by_case.setdefault(case, []).append(stem(p))
        rows.append({"file": os.path.basename(p), "case_number": case, "file_date": filed,
                     "pages": body.count("=== PAGE "), "status": status})
    for r in rows:
        if r["case_number"] and len(by_case[r["case_number"]]) > 1 and r["status"] == "new":
            r["status"] = "new (DUPLICATE in folder)"

    with open(os.path.join(OUT_DIR, "summary.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)

    print(f"\nOCR done in {time.time() - start:.0f}s. Text in {OUT_DIR}\n")
    for r in sorted(rows, key=lambda r: (r["file_date"], r["case_number"])):
        print(f"  {r['file_date'] or '??????????'}  {r['case_number'] or '?':<16} "
              f"{r['pages']:>3}p  {r['status']:<26} {r['file']}")
    tally = {}
    for r in rows:
        tally[r["status"]] = tally.get(r["status"], 0) + 1
    print("\n  " + "  |  ".join(f"{k}: {v}" for k, v in sorted(tally.items())))
    return 0


if __name__ == "__main__":
    sys.exit(main())
