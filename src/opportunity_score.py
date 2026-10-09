"""Rank a batch's new records as buying opportunities (Jeff, 2026-10-09).

Builds the ranked table Jeff signed off on for the 2026-10-08 probate run
(`# | Property | ZIP (rank) | Value / equity | Owned | Heirs | Reach | Why`) and
adds a 0-100 Score column. The table and its columns are the fallback: if the
score ever reads wrong, the same facts are on the row to judge by hand.

    python src/opportunity_score.py --batch "output/probate_batch_chunk*_20261008.xlsx" \
        --batch output/probate_batch.xlsx --type probate
    python src/opportunity_score.py --batch output/petition_batch.xlsx --type foreclosure

Read-only and free: DataSift reads of the batch's OWN records (never an account
sweep) plus the Tulsa Assessor's deed history. Writes output/opportunity_<type>_<date>.md
and a .json beside it.

Score (100):
  situation 30   probate: heir count, no living spouse; foreclosure: borrowers,
                 owner alive, modification-exhausted
  zip       20   rank on the DataSift county playbook "Best ZIP codes" list
  equity    15   DataSift equity %, or (value - unpaid balance) for foreclosure
  owned     10   years in the family, from the county deed chain
  reach     10   Dial First x2 + Dial Second, capped
  vacancy   10   absentee owner list, or nobody left living in the house
  value      5   inside a normal buy box
Caps (score held under 50, i.e. weak, whatever the points say): a living spouse
in a probate, a contested will, foreclosure debt above 90% of value.
"""

import argparse
import datetime as dt
import glob
import json
import os
import re
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# learn.datasift.ai/county-list-playbook, Tulsa County "Best ZIP codes", as shown
# 2026-10-09 (top 18 of the list). Update when the page changes; a ZIP not listed
# scores as "not in top 18".
ZIP_RANKS = {
    "74012": 1, "74114": 2, "74107": 3, "74105": 4, "74106": 5, "74134": 6,
    "74104": 7, "74133": 8, "74135": 9, "74008": 10, "74055": 11, "74011": 12,
    "74037": 13, "74136": 14, "74132": 15, "74115": 16, "74126": 17, "74063": 18,
}
ZIP_LIST_LEN = 18

SPOUSE_RX = re.compile(r"(?<!no )(?<!not )surviving\s+(wife|husband|spouse)", re.I)
PAPERWORK_DEEDS = ("QUIT CLAIM", "SURVIVING", "TRANSFER ON DEATH", "TOD", "TRUSTEE")
NAME_ONLY_PENALTY = 8
SPOUSE_PR = {"wife", "husband", "spouse"}
TIERS = ((75, "Fantastic"), (50, "Good"), (0, "Weak"))


def _num(v):
    try:
        return float(str(v).replace("$", "").replace(",", ""))
    except (TypeError, ValueError):
        return None


def _year(datestr):
    m = re.search(r"(\d{4})", datestr or "")
    return int(m.group(1)) if m else None


def owned_since(deeds, title_holder, today=None):
    """Year the family took title, walking the deed chain newest to oldest.

    Paperwork inside the family does not restart the clock: a $0 quit claim,
    survivorship affidavit, transfer-on-death or trustee deed, or any $0 deed
    between people with the title holder's surname. The first deed that is a
    real purchase (a price, or a $0 "History" deed from someone outside the
    family) is when the family got the house. Returns (year, label); year None
    when the chain never leaves the family ("before the 1980s").
    """
    surname = (title_holder or "").split(",")[0].strip().upper()
    for d in deeds or []:  # newest first
        grantee = (d.get("grantee") or "").upper()
        grantor = (d.get("grantor") or "").upper()
        yr = _year(d.get("sale_date"))
        if not grantee and not grantor:  # "History - Unknown": oldest date on file
            return yr, str(yr)
        price = _num(d.get("sale_price")) or 0
        dtype = (d.get("deed_type") or "").upper()
        family = bool(surname) and surname in grantee and surname in grantor
        paperwork = any(k in dtype for k in PAPERWORK_DEEDS)
        if price == 0 and (family or paperwork):
            continue
        return yr, str(yr)
    return None, "before the 1980s"


def _living_spouse(sheet):
    if ";" in str(sheet.get("Date of Death") or ""):
        return False  # two decedents (both parents)
    if SPOUSE_RX.search(str(sheet.get("Marital Status") or "")):
        return True
    return str(sheet.get("PR Relationship") or "").strip().lower() in SPOUSE_PR


def score_record(rec, kind="probate", today=None):
    """rec: the property fields + 'sheet' (the batch row) + 'deeds' + DF/DS."""
    today = today or dt.date.today()
    s = rec.get("sheet") or {}
    parts, caps, flags, why = {}, [], [], []
    lists = rec.get("lists") or []

    # situation
    if kind == "probate":
        heirs = int(_num(s.get("Heir Count")) or 0)
        base = {0: 15, 1: 30, 2: 25, 3: 20, 4: 15}.get(heirs, 10 if heirs < 7 else 5)
        spouse = _living_spouse(s)
        if spouse:
            caps.append("living spouse")
            why.append("a spouse is still living and likely stays")
        else:
            why.append(f"{heirs} heir{'s' if heirs != 1 else ''}, no spouse in the house")
        if "CONTESTED" in str(s.get("Testate") or "").upper():
            caps.append("contested will")
        parts["situation"] = base
    else:
        alive = str(s.get("Owner Alive") or "Yes").strip().lower() != "no"
        cob = bool(str(s.get("Co-Borrower Last Name") or "").strip())
        base = (20 if cob else 25) if alive else 5
        mods = int(_num(s.get("Loan Modification Count")) or 0)
        if mods >= 2:
            base += 5
            why.append(f"{mods} loan modifications")
        parts["situation"] = min(base, 30)
        spouse = False
        if not alive:
            why.append("owner deceased, no living contact named")

    # zip
    z = str(((rec.get("sheet") or {}).get("Property Zip")) or rec.get("zip") or "")[:5]
    rank = ZIP_RANKS.get(z)
    parts["zip"] = round(20 - (rank - 1) * (13 / (ZIP_LIST_LEN - 1)), 1) if rank else 4
    rec["zip_label"] = f"{z} (#{rank})" if rank else f"{z} (not in top {ZIP_LIST_LEN})"

    # equity
    value = _num(rec.get("estimate_value"))
    eq = _num(rec.get("equity_percent"))
    if kind == "foreclosure":
        upb = _num(s.get("Unpaid Principal Balance"))
        if value and upb is not None:
            eq = max(0.0, (value - upb) / value * 100)
            if upb > 0.9 * value:
                caps.append("debt over 90% of value")
    if eq is None:
        eq = 100.0 if "Free & Clear" in lists else 50.0
        if "Free & Clear" not in lists:
            flags.append("equity unknown")
    parts["equity"] = round(15 * min(eq, 100) / 100, 1)
    rec["equity_label"] = "F&C" if eq >= 99.5 else f"{eq:.0f}%"

    # owned
    if rec.get("deeds") is None:  # no county parcel id: history not looked up
        held, parts["owned"], rec["owned_label"] = 0, 5, "Unknown"
        flags.append("no deed history")
    else:
        yr, note = owned_since(rec["deeds"], s.get("Title Holder of Record"))
        held = (today.year - yr) if yr else 45
        parts["owned"] = 10 if held >= 30 else 8 if held >= 20 else 6 if held >= 10 else 3 if held >= 5 else 1
        rec["owned_label"] = f"Since {note}" if yr else "Before the 1980s"
    if held >= 20:
        why.append("long held, likely dated")

    # reach
    df, ds = int(rec.get("DF") or 0), int(rec.get("DS") or 0)
    parts["reach"] = min(10, df * 2 + ds)
    if df + ds == 0:
        flags.append("0 good numbers: deep prospect or mail")

    # vacancy
    if "Absentee Owners" in lists:
        parts["vacancy"] = 10
        why.append("absentee")
    elif kind == "probate" and not spouse:
        parts["vacancy"] = 6
    else:
        parts["vacancy"] = 0

    # value fit
    if value is None:
        parts["value"] = 0
        flags.append("no property data in DataSift")
    elif value > 700_000:
        parts["value"] = 0
        flags.append("far above buy box")
    else:
        parts["value"] = 5 if 60_000 <= value <= 450_000 else 2

    title = str(s.get("Title Holder of Record") or "").upper()
    if "NAME-ONLY" in title or "INITIALS-ONLY" in title:
        flags.append("name-only property match")
        parts["name_only"] = -NAME_ONLY_PENALTY

    score = round(sum(parts.values()))
    if caps:
        score = min(score, 49)
    tier = next(t for floor, t in TIERS if score >= floor)
    return {"score": score, "tier": tier, "parts": parts, "caps": caps,
            "flags": flags, "why": "; ".join(why)}


def gather(batch_paths, kind):
    """Read the batch rows, then ONLY those records from DataSift + the Assessor."""
    import datasift_api as D
    from openpyxl import load_workbook
    from tulsa_assessor import get_parcel_sales_history

    rows = {}
    for path in batch_paths:
        ws = load_workbook(path, read_only=True).active
        it = ws.iter_rows(values_only=True)
        hdr = next(it)
        for r in it:
            d = dict(zip(hdr, r))
            if d.get("Property Street"):
                rows[(d["Property Street"], d.get("Property City"))] = d
    out = []
    for (street, city), sheet in rows.items():
        hit = D.find_property_by_address(street, city or "", "OK")
        if not hit:
            continue
        p = D.get_property(hit["uuid"])
        phones = D.get_owner(p["owner"]["uuid"]).get("phones") or []
        if not phones and not D.get_message_board(p["owner"]["uuid"]):
            continue  # not created by this run (held or excluded rows)
        rec = {k: p.get(k) for k in ("uuid", "estimate_value", "equity_percent", "sqft",
                                     "year", "rental_value", "mls", "lists", "neighborhood")}
        rec.update(street=street, city=city, sheet={k: v for k, v in sheet.items() if v not in (None, "")})
        rec["DF"] = sum(1 for x in phones if "Dial First" in (x.get("tags") or []))
        rec["DS"] = sum(1 for x in phones if "Dial Second" in (x.get("tags") or []))
        pid = sheet.get("Parcel ID")
        rec["zip"] = str((p.get("address") or {}).get("postal_code") or "")[:5]
        rec["deeds"] = [] if pid else None
        if pid:
            for attempt in range(3):
                try:
                    rec["deeds"] = get_parcel_sales_history(pid)
                    break
                except Exception:
                    time.sleep(8)
            time.sleep(3)
        out.append(rec)
    return out


def _money(v):
    v = _num(v)
    return f"${v / 1000:.0f}K" if v else "Unknown"


def render(recs, kind):
    head = ("| # | Property | Score | ZIP (rank) | Value / equity | Owned | "
            + ("Heirs" if kind == "probate" else "Borrowers") + " | Reach | Why |")
    lines = [head, "|" + "---|" * 9]
    for i, r in enumerate(recs, 1):
        sc, s = r["score_detail"], r.get("sheet") or {}
        if kind == "probate":
            n = int(_num(s.get('Heir Count')) or 0)
            who = f"{n} heir{'s' if n != 1 else ''}, PR {s.get('PR Relationship') or '?'}"
        else:
            who = "2 (co-borrower)" if s.get("Co-Borrower Last Name") else "1"
        reach = f"DF {r['DF']}" + (f", DS {r['DS']}" if r["DS"] else "") if r["DF"] + r["DS"] else "**0 phones**"
        why = sc["why"]
        if sc["caps"]:
            why += "; held weak: " + ", ".join(sc["caps"])
        if sc["flags"]:
            why += "; flag: " + ", ".join(sc["flags"])
        lines.append(f"| {i} | **{r['street']}**, {r.get('city') or ''} | **{sc['score']}** {sc['tier']} | "
                     f"{r['zip_label']} | {_money(r.get('estimate_value'))}, {r['equity_label']} | "
                     f"{r['owned_label']} | {who} | {reach} | {why} |")
    return "\n".join(lines)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--batch", action="append", default=[], help="batch sheet path or glob (repeatable)")
    ap.add_argument("--type", choices=("probate", "foreclosure"), default="probate")
    ap.add_argument("--from-json", help="re-score a saved gather (no network)")
    args = ap.parse_args()

    if args.from_json:
        recs = json.load(open(args.from_json, encoding="utf-8"))
    else:
        paths = sorted({p for g in args.batch for p in glob.glob(g)})
        if not paths:
            sys.exit("no batch sheets matched")
        recs = gather(paths, args.type)
    if not recs:
        sys.exit("no records found for this batch (check the sheets and that the run created them)")
    for r in recs:
        r["score_detail"] = score_record(r, args.type)
    recs.sort(key=lambda r: -r["score_detail"]["score"])

    stamp = dt.date.today().isoformat()
    base = os.path.join(ROOT, "output", f"opportunity_{args.type}_{stamp}")
    table = render(recs, args.type)
    with open(base + ".md", "w", encoding="utf-8") as f:
        f.write(table + "\n")
    with open(base + ".json", "w", encoding="utf-8") as f:
        json.dump(recs, f, default=str, indent=1)
    print(table)
    print(f"\n{len(recs)} records -> {base}.md")


if __name__ == "__main__":
    main()
